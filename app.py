import ipaddress
import io
import logging
import os
import re
import shutil
import socket
import subprocess
import sys
import time
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

from flask import Flask, g, jsonify, render_template, request, send_file
import psutil
from werkzeug.middleware.proxy_fix import ProxyFix


BASE_DIR = Path(__file__).resolve().parent
MAX_NETWORKS = int(os.getenv("MAX_NETWORKS", "10"))
MAX_ADDRESSES_PER_NETWORK = int(os.getenv("MAX_ADDRESSES_PER_NETWORK", "4096"))
SCAN_TIMEOUT = int(os.getenv("SCAN_TIMEOUT", "120"))
TRUST_PROXY = os.getenv("TRUST_PROXY", "false").lower() in {"1", "true", "yes", "on"}

app = Flask(__name__, template_folder=BASE_DIR / "templates", static_folder=BASE_DIR / "static")
if TRUST_PROXY:
    # Only enable this when requests always pass through one trusted reverse proxy.
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)

access_logger = logging.getLogger("netscope.access")
access_logger.setLevel(getattr(logging, os.getenv("LOG_LEVEL", "INFO").upper(), logging.INFO))
if not access_logger.handlers:
    access_handler = logging.StreamHandler(sys.stdout)
    access_handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s", "%Y-%m-%d %H:%M:%S"))
    access_logger.addHandler(access_handler)
access_logger.propagate = False


@app.before_request
def begin_request_log():
    g.request_started_at = time.perf_counter()


@app.after_request
def write_access_log(response):
    elapsed_ms = round((time.perf_counter() - getattr(g, "request_started_at", time.perf_counter())) * 1000)
    access_logger.info(
        "access ip=%s method=%s path=%s status=%s duration_ms=%s",
        request.remote_addr or "-",
        request.method,
        request.path,
        response.status_code,
        elapsed_ms,
    )
    return response


def resolve_nmap():
    configured = os.getenv("NMAP_PATH")
    if configured and Path(configured).is_file():
        return configured

    discovered = shutil.which("nmap")
    if discovered:
        return discovered

    if os.name == "nt":
        candidates = [
            Path(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")) / "Nmap" / "nmap.exe",
            Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "Nmap" / "nmap.exe",
        ]
        for candidate in candidates:
            if candidate.is_file():
                return str(candidate)
    return None


def normalize_network(value):
    if not isinstance(value, str) or not value.strip():
        raise ValueError("网段不能为空")

    value = value.strip()
    try:
        network = ipaddress.ip_network(value, strict=False)
    except ValueError as exc:
        raise ValueError(f"无效网段：{value}") from exc

    if network.version != 4:
        raise ValueError("当前仅支持 IPv4 网段")
    if network.num_addresses > MAX_ADDRESSES_PER_NETWORK:
        raise ValueError(
            f"网段 {network.with_prefixlen} 包含 {network.num_addresses} 个地址，"
            f"超过单网段上限 {MAX_ADDRESSES_PER_NETWORK}"
        )
    return network.with_prefixlen


def parse_nmap_output(output):
    hosts = []
    current = None

    for raw_line in output.splitlines():
        line = raw_line.strip()
        if line.startswith("Nmap scan report for "):
            target = line.removeprefix("Nmap scan report for ")
            match = re.match(r"^(.*?)\s*\(([^)]+)\)$", target)
            if match:
                current = {"hostname": match.group(1).strip(), "ip": match.group(2).strip()}
            else:
                current = {"hostname": "", "ip": target.strip()}
            hosts.append(current)
        elif current and line.startswith("MAC Address: "):
            mac_info = line.removeprefix("MAC Address: ")
            match = re.match(r"^([0-9A-Fa-f:]{17})(?:\s+\((.*)\))?$", mac_info)
            if match:
                current["mac"] = match.group(1).upper()
                current["vendor"] = (match.group(2) or "").strip()

    return hosts


def normalize_mac(value):
    if not value:
        return None
    compact = re.sub(r"[^0-9A-Fa-f]", "", value)
    if len(compact) != 12 or set(compact) == {"0"}:
        return None
    return ":".join(compact[index:index + 2] for index in range(0, 12, 2)).upper()


def get_primary_network_info():
    """Return primary IPv4 interface details on Windows and Linux."""
    primary_ip = None
    probe = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        # UDP connect does not send traffic; it only asks the OS which route it would use.
        probe.connect(("8.8.8.8", 80))
        primary_ip = probe.getsockname()[0]
    except OSError:
        pass
    finally:
        probe.close()

    candidates = []
    for interface, addresses in psutil.net_if_addrs().items():
        mac = next(
            (normalize_mac(address.address) for address in addresses if address.family == psutil.AF_LINK),
            None,
        )
        for address in addresses:
            if address.family != socket.AF_INET or not address.netmask:
                continue
            try:
                ip = ipaddress.ip_address(address.address)
                network = ipaddress.ip_network(f"{address.address}/{address.netmask}", strict=False)
            except ValueError:
                continue
            if ip.is_loopback or not ip.is_private:
                continue
            info = {
                "ip": address.address,
                "network": network.with_prefixlen,
                "netmask": address.netmask,
                "mac": mac,
                "interface": interface,
            }
            if address.address == primary_ip:
                return info
            candidates.append(info)

    return candidates[0] if candidates else None


def get_primary_network():
    info = get_primary_network_info()
    return info["network"] if info else None


def scan_network(network):
    command = [resolve_nmap() or "nmap", "-sn", network]
    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=SCAN_TIMEOUT,
            check=False,
        )
    except subprocess.TimeoutExpired:
        return {"network": network, "hosts": [], "status": "error", "error": f"扫描超时（{SCAN_TIMEOUT} 秒）"}
    except OSError as exc:
        return {"network": network, "hosts": [], "status": "error", "error": f"无法执行 nmap：{exc}"}

    if completed.returncode != 0:
        detail = (completed.stderr or completed.stdout or "nmap 执行失败").strip()
        return {"network": network, "hosts": [], "status": "error", "error": detail[-600:]}

    hosts = parse_nmap_output(completed.stdout)
    local_info = get_primary_network_info()
    if local_info:
        for host in hosts:
            if host["ip"] == local_info["ip"]:
                if local_info.get("mac") and not host.get("mac"):
                    host["mac"] = local_info["mac"]
                if not host.get("vendor"):
                    host["vendor"] = "本机网卡"
                if not host.get("hostname"):
                    host["hostname"] = socket.gethostname()
    return {"network": network, "hosts": hosts, "count": len(hosts), "status": "success"}


def prepare_export_data(payload):
    raw_networks = payload.get("networks", [])
    raw_results = payload.get("results", [])
    if not isinstance(raw_networks, list) or not isinstance(raw_results, list):
        raise ValueError("导出数据格式错误")
    if not raw_networks or len(raw_networks) > MAX_NETWORKS:
        raise ValueError("没有可导出的检测网段")

    networks = [normalize_network(value) for value in raw_networks]
    groups = []
    total_hosts = 0
    for result in raw_results:
        if not isinstance(result, dict):
            continue
        network = str(result.get("network", ""))[:80]
        hosts = []
        for host in result.get("hosts", []):
            if not isinstance(host, dict) or total_hosts >= MAX_ADDRESSES_PER_NETWORK * MAX_NETWORKS:
                continue
            hosts.append({
                "ip": str(host.get("ip", ""))[:64],
                "hostname": str(host.get("hostname", ""))[:200] or "-",
                "mac": str(host.get("mac", ""))[:64] or "-",
                "vendor": str(host.get("vendor", ""))[:200] or "-",
            })
            total_hosts += 1
        groups.append({"network": network, "hosts": hosts})
    return networks, groups, str(payload.get("search", ""))[:64]


def export_text(networks, groups, search, markdown=False):
    lines = ["# 网段 IP 检测结果" if markdown else "网段 IP 检测结果"]
    lines.append(f"检测网段：{', '.join(networks)}")
    if search:
        lines.append(f"IP 筛选：{search}")
    lines.append("")
    for group in groups:
        lines.append(f"## {group['network']}" if markdown else f"网段：{group['network']}")
        if markdown:
            lines.extend(["", "| 序号 | IP 地址 | 主机名 | MAC 地址 | 设备厂商 | 状态 |", "|---:|---|---|---|---|---|"])
            for index, host in enumerate(group["hosts"], 1):
                values = [str(index), host["ip"], host["hostname"], host["mac"], host["vendor"], "在线"]
                lines.append("| " + " | ".join(value.replace("|", "\\|") for value in values) + " |")
        else:
            lines.append("序号\tIP 地址\t主机名\tMAC 地址\t设备厂商\t状态")
            for index, host in enumerate(group["hosts"], 1):
                lines.append(f"{index}\t{host['ip']}\t{host['hostname']}\t{host['mac']}\t{host['vendor']}\t在线")
        if not group["hosts"]:
            lines.append("暂无匹配的在线设备")
        lines.append("")
    return "\n".join(lines).encode("utf-8-sig")


def export_xlsx(networks, groups, search):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "检测结果"
    sheet.merge_cells("A1:F1")
    sheet["A1"] = f"检测网段：{', '.join(networks)}"
    sheet["A1"].font = Font(bold=True, size=13)
    row = 2
    if search:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        sheet.cell(row, 1, f"IP 筛选：{search}")
        row += 1
    headers = ["序号", "IP 地址", "主机名", "MAC 地址", "设备厂商", "状态"]
    for group in groups:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = sheet.cell(row, 1, f"网段：{group['network']}")
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F2F6FC")
        row += 1
        for column, header in enumerate(headers, 1):
            cell = sheet.cell(row, column, header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        row += 1
        if group["hosts"]:
            for index, host in enumerate(group["hosts"], 1):
                values = [index, host["ip"], host["hostname"], host["mac"], host["vendor"], "在线"]
                for column, value in enumerate(values, 1):
                    sheet.cell(row, column, value)
                row += 1
        else:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            sheet.cell(row, 1, "暂无匹配的在线设备").alignment = Alignment(horizontal="center")
            row += 1
        row += 1
    for letter, width in {"A": 9, "B": 18, "C": 25, "D": 20, "E": 32, "F": 10}.items():
        sheet.column_dimensions[letter].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_docx(networks, groups, search):
    from docx import Document

    document = Document()
    document.add_heading("网段 IP 检测结果", level=0)
    document.add_paragraph(f"检测网段：{', '.join(networks)}")
    if search:
        document.add_paragraph(f"IP 筛选：{search}")
    headers = ["序号", "IP 地址", "主机名", "MAC 地址", "设备厂商", "状态"]
    for group in groups:
        document.add_heading(group["network"], level=1)
        table = document.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        for cell, header in zip(table.rows[0].cells, headers):
            cell.text = header
        if group["hosts"]:
            for index, host in enumerate(group["hosts"], 1):
                values = [index, host["ip"], host["hostname"], host["mac"], host["vendor"], "在线"]
                for cell, value in zip(table.add_row().cells, values):
                    cell.text = str(value)
        else:
            table.add_row().cells[0].text = "暂无匹配的在线设备"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def export_pdf(networks, groups, search):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.enums import TA_LEFT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    for style_name in ("Title", "Heading2", "BodyText"):
        styles[style_name].fontName = "STSong-Light"
        styles[style_name].alignment = TA_LEFT
    story = [Paragraph("网段 IP 检测结果", styles["Title"]), Paragraph(f"检测网段：{', '.join(networks)}", styles["BodyText"])]
    if search:
        story.append(Paragraph(f"IP 筛选：{search}", styles["BodyText"]))
    story.append(Spacer(1, 12))
    headers = ["序号", "IP 地址", "主机名", "MAC 地址", "设备厂商", "状态"]
    for group in groups:
        story.append(Paragraph(group["network"], styles["Heading2"]))
        rows = [headers]
        for index, host in enumerate(group["hosts"], 1):
            rows.append([index, host["ip"], host["hostname"], host["mac"], host["vendor"], "在线"])
        if not group["hosts"]:
            rows.append(["", "暂无匹配的在线设备", "", "", "", ""])
        table = Table(rows, colWidths=[42, 95, 125, 115, 220, 55], repeatRows=1)
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F6FC")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DCDFe6")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (0, -1), "CENTER"),
            ("PADDING", (0, 0), (-1, -1), 6),
        ]))
        story.extend([table, Spacer(1, 14)])
    document.build(story)
    return output.getvalue()


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/health")
def health():
    nmap_binary = resolve_nmap()
    return jsonify({"ok": True, "nmap": bool(nmap_binary), "nmap_path": nmap_binary})


@app.get("/api/local-network")
def local_network():
    info = get_primary_network_info()
    if not info:
        return jsonify({"error": "无法识别当前 IPv4 网段"}), 404
    return jsonify(info)


@app.post("/api/scan")
def scan():
    payload = request.get_json(silent=True) or {}
    raw_networks = payload.get("networks")
    if not isinstance(raw_networks, list):
        return jsonify({"error": "networks 必须是网段数组"}), 400

    try:
        networks = list(dict.fromkeys(normalize_network(value) for value in raw_networks))
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    if not networks:
        return jsonify({"error": "请至少输入一个网段"}), 400
    if len(networks) > MAX_NETWORKS:
        return jsonify({"error": f"一次最多扫描 {MAX_NETWORKS} 个网段"}), 400
    if not resolve_nmap():
        return jsonify({"error": "服务器未安装 nmap，请先安装后再试"}), 503

    started_at = time.perf_counter()
    results_by_network = {}
    with ThreadPoolExecutor(max_workers=min(len(networks), 4)) as executor:
        futures = {executor.submit(scan_network, network): network for network in networks}
        for future in as_completed(futures):
            network = futures[future]
            try:
                results_by_network[network] = future.result()
            except Exception:
                app.logger.exception("扫描网段时出现异常：%s", network)
                results_by_network[network] = {
                    "network": network,
                    "hosts": [],
                    "status": "error",
                    "error": "扫描时发生内部错误",
                }

    elapsed_ms = round((time.perf_counter() - started_at) * 1000)
    return jsonify({
        "results": [results_by_network[network] for network in networks],
        "elapsed_ms": elapsed_ms,
    })


@app.post("/api/export")
def export_results():
    payload = request.get_json(silent=True) or {}
    export_format = str(payload.get("format", "")).lower()
    try:
        networks, groups, search = prepare_export_data(payload)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    exporters = {
        "xlsx": (export_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
        "docx": (export_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
        "pdf": (export_pdf, "application/pdf", "pdf"),
        "md": (lambda n, g, s: export_text(n, g, s, markdown=True), "text/markdown; charset=utf-8", "md"),
        "txt": (export_text, "text/plain; charset=utf-8", "txt"),
    }
    if export_format not in exporters:
        return jsonify({"error": "不支持的导出格式"}), 400
    exporter, mimetype, extension = exporters[export_format]
    try:
        content = exporter(networks, groups, search)
    except Exception:
        app.logger.exception("生成导出文件失败")
        return jsonify({"error": "生成导出文件失败，请检查相关依赖"}), 500
    filename = f"ip-scan-{datetime.now().strftime('%Y%m%d-%H%M%S')}.{extension}"
    return send_file(io.BytesIO(content), mimetype=mimetype, as_attachment=True, download_name=filename)


if __name__ == "__main__":
    host = os.getenv("HOST", "0.0.0.0")
    port = int(os.getenv("PORT", "5000"))
    access_logger.info("service_start host=%s port=%s nmap=%s", host, port, resolve_nmap() or "not-found")
    app.run(host=host, port=port, debug=False)
